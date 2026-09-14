import os
import sys
import json
import openpyxl
# -- coding: utf-8 --**


# ui(scui_ui_res) 基准: GUI 注入, 日志路径相对化显示
SCUI_UI_ROOT = None
# tools 基准: GUI 注入(读句柄偏移配置)
SCUI_TOOLS = None


#============================================================
# 表布局(固化书写, 原 scui_res_lang.json 字段已并入此处)
#   第1行   : A1=编号区标记, B1=拼音区标记, C1起=语言列表
#   第1列   : 编号区 ID(SCUI_LANG_IDX_0Xxxxx, excel 表达式输出)
#   第2列   : 拼音区 ID(SCUI_LANG_ABBR_xxxx, 自定义输出)
#   第3列起 : 各语言文本(顺序与第1行语言列表一一对应)
#============================================================
SCUI_LANG_XLSX   = 'scui_lang_parser.xlsx'
SCUI_LANG_SHEET  = 'translate'
# 固化默认语言列表(仅旧布局迁移时使用; 新布局以第1行为准)
SCUI_LANG_LANG_DEF = ['zh', 'en', 'de', 'fra', 'nl', 'pt', 'jp']
# 固化 scui_lang_type_t 的固定项(原 scui_res_lang.json 的 custom 字段)
SCUI_LANG_CUSTOM = ['multi', 'ascii', 'symbol']
# 固化句柄偏移默认(前端"设置句柄"弹窗可覆盖)
SCUI_LANG_OFFSET_NAME  = 'SCUI_HANDLE_OFFSET_LANG'
SCUI_LANG_OFFSET_VALUE = '0x5000 - 1'
# 表定位(0基)
SCUI_LANG_ROW_HEAD  = 0             # 第1行: 语言列表
SCUI_LANG_COL_IDX   = 0             # 第1列: 编号区
SCUI_LANG_COL_ABBR  = 1             # 第2列: 拼音区
SCUI_LANG_COL_DATA  = 2             # 第3列起: 语言数据
SCUI_LANG_HEAD_IDX  = '编号区'
SCUI_LANG_HEAD_ABBR = '拼音区'
# 拼音区: 最多取前 N 个汉字的首字母(可能不足 N 个)
SCUI_LANG_ABBR_MAX  = 6
# 归一化统一写入的默认单元格属性列宽
SCUI_LANG_COL_W_IDX  = 20           # 编号区列宽
SCUI_LANG_COL_W_ABBR = 26           # 拼音区列宽
SCUI_LANG_COL_W_TEXT = 40           # 语言文本列宽


# 句柄偏移: 前端"设置句柄"写入的 json 覆盖, 无则用配置默认
def _handle_offset(tag, name, value):
    try:
        cfg = os.path.join(SCUI_TOOLS or os.path.dirname(__file__), 'scui_pack_handle.json')
        if os.path.isfile(cfg):
            j = json.load(open(cfg, encoding='utf-8'))
            item = j.get(tag, {})
            if item.get('name'):
                name = item['name']
            if item.get('value'):
                value = item['value']
    except Exception:
        pass
    return name, value


# 日志路径显示: 相对 ui(scui_ui_res) 打印, 命令行(GUI未注入)保持绝对
def _rel_ui(p):
    if not SCUI_UI_ROOT:
        return p
    try:
        r = os.path.relpath(p, SCUI_UI_ROOT).replace('\\', '/')
        if r.startswith('..'):
            return p
        return 'scui_ui_res => %s' % r
    except Exception:
        return p


#============================================================
# 拼音区 ID 生成(自定义方式): 最多取前 N 个汉字的首字母
#============================================================
_PYPINYIN = None


def _pinyin_initials(text):
    global _PYPINYIN
    if _PYPINYIN is None:
        try:
            from pypinyin import lazy_pinyin, Style
            _PYPINYIN = (lazy_pinyin, Style.FIRST_LETTER)
        except Exception:
            _PYPINYIN = False
    if not _PYPINYIN:
        return ''
    lazy_pinyin, style = _PYPINYIN
    try:
        return ''.join(lazy_pinyin(text, style=style))
    except Exception:
        return ''


def _is_han(c):
    return ('\u4e00' <= c <= '\u9fff') or ('\u3400' <= c <= '\u4dbf')


def _ascii_alnum(c):
    return ('0' <= c <= '9') or ('A' <= c <= 'Z')


def scui_lang_abbr(text):
    """拼音区后缀: 最多前 N 个汉字的首字母(大写)
    退化: 无汉字时取文本中 ASCII 字母数字前 N 个; 仍为空则返回空串(由调用方兜底)
    """
    han = ''.join([c for c in text if _is_han(c)][:SCUI_LANG_ABBR_MAX])
    s = _pinyin_initials(han) if han else ''
    s = ''.join([c for c in s.upper() if _ascii_alnum(c)])
    if not s:
        s = ''.join([c for c in text.upper() if _ascii_alnum(c)])[:SCUI_LANG_ABBR_MAX]
    return s


def _abbr_unique(name, used):
    """拼音区重复时追加 _R1~n 尾缀"""
    if name not in used:
        used.add(name)
        return name
    idx = 1
    while '%s_R%d' % (name, idx) in used:
        idx += 1
    name = '%s_R%d' % (name, idx)
    used.add(name)
    return name


#============================================================
# 工作表读写(新布局)
#============================================================
def _cell_text(cell):
    v = cell.value
    return '' if v is None else str(v)


def _sheet_path(src_path):
    return os.path.join(src_path, SCUI_LANG_XLSX)


def _sheet_pick(wb):
    if SCUI_LANG_SHEET in wb.sheetnames:
        return wb[SCUI_LANG_SHEET]
    return wb[wb.sheetnames[0]]


def _sheet_layout(ws):
    """读取: 语言列表 + 数据行/列范围(1基, 含端点)"""
    langs = []
    for c in range(SCUI_LANG_COL_DATA + 1, ws.max_column + 1):
        langs.append(_cell_text(ws.cell(SCUI_LANG_ROW_HEAD + 1, c)).strip())
    while langs and langs[-1] == '':
        langs.pop()
    row_s = SCUI_LANG_ROW_HEAD + 2
    row_e = ws.max_row
    col_hint = SCUI_LANG_COL_DATA + 1
    while row_e > row_s and _cell_text(ws.cell(row_e, col_hint)) == '':
        row_e -= 1
    col_s = SCUI_LANG_COL_DATA + 1
    col_e = col_s + len(langs) - 1
    return langs, row_s, row_e, col_s, col_e


def _legacy_migrate(ws, log):
    """旧布局(顶格, 无表头) -> 新布局: 顶行插入语言行, 数据整体右移两列"""
    max_r, max_c = ws.max_row, ws.max_column
    log('lang 检测到旧布局(顶格), 迁移为: 第1行语言表头 + 首两列ID')
    data = [[ws.cell(r, c).value for c in range(1, max_c + 1)] for r in range(1, max_r + 1)]
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            ws.cell(r, c).value = None
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            ws.cell(r + 1, c + 2).value = data[r - 1][c - 1]
    ws.cell(1, SCUI_LANG_COL_IDX + 1).value = SCUI_LANG_HEAD_IDX
    ws.cell(1, SCUI_LANG_COL_ABBR + 1).value = SCUI_LANG_HEAD_ABBR
    for i, lang in enumerate(SCUI_LANG_LANG_DEF[:max_c]):
        ws.cell(1, SCUI_LANG_COL_DATA + 1 + i).value = lang


def normalize_scui_lang_sheet(src_path, log=print):
    """归一化: 第1列写 excel 表达式(编号区), 第2列写自定义文本(拼音区); 直接复写 xlsx
    返回 (语言列表, 数据行数)
    """
    path = _sheet_path(src_path)
    wb = openpyxl.load_workbook(path, read_only=False)
    ws = _sheet_pick(wb)

    # 旧布局自动迁移
    if _cell_text(ws.cell(1, SCUI_LANG_COL_IDX + 1)).strip() != SCUI_LANG_HEAD_IDX:
        _legacy_migrate(ws, log)

    langs, row_s, row_e, col_s, col_e = _sheet_layout(ws)
    if not langs:
        log('lang 语言列表为空(第1行), 跳过归一化')
        wb.close()
        return [], 0

    # 编号区: excel 表达式(行号 - 首个数据行号)
    expr = '="SCUI_LANG_IDX_0X"&LOWER(TEXT(ROW()-{0},"0000"))'.format(row_s)
    used = set()
    for r in range(row_s, row_e + 1):
        ws.cell(r, SCUI_LANG_COL_IDX + 1).value = expr

        # 拼音区: 最多前 6 个汉字首字母; 重复追加 _R1~n
        zh = _cell_text(ws.cell(r, col_s))
        suffix = scui_lang_abbr(zh)
        if not suffix:
            suffix = '0X%04X' % (r - row_s)
        full = _abbr_unique('SCUI_LANG_ABBR_' + suffix, used)
        ws.cell(r, SCUI_LANG_COL_ABBR + 1).value = full

    # 默认单元格属性: 表头与编号区居中, 拼音区左对齐, 语言文本左对齐+自动换行(完全显示)
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    align_ctr = Alignment(horizontal='center', vertical='center', wrap_text=False)
    align_lft = Alignment(horizontal='left',   vertical='center', wrap_text=False)
    align_wrp = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    for c in range(1, col_e + 1):
        ws.cell(SCUI_LANG_ROW_HEAD + 1, c).alignment = align_ctr

    for r in range(row_s, row_e + 1):
        ws.cell(r, SCUI_LANG_COL_IDX + 1).alignment = align_ctr
        ws.cell(r, SCUI_LANG_COL_ABBR + 1).alignment = align_lft
        for ci in range(len(langs)):
            ws.cell(r, col_s + ci).alignment = align_wrp

    # 行高交还 Excel 自适应(清除手工固定高度), 换行内容才能完全显示
    for r in range(1, row_e + 1):
        if r in ws.row_dimensions:
            del ws.row_dimensions[r]

    # 列宽: 归一化统一写入(与对齐/换行同属默认属性)
    col_wide = {SCUI_LANG_COL_IDX + 1: SCUI_LANG_COL_W_IDX,
                SCUI_LANG_COL_ABBR + 1: SCUI_LANG_COL_W_ABBR}
    for ci in range(len(langs)):
        col_wide[col_s + ci] = SCUI_LANG_COL_W_TEXT
    for col, wide in col_wide.items():
        ws.column_dimensions[get_column_letter(col)].width = wide
    log('lang 默认属性: 对齐/换行已写入, 行高转自适应, 列宽已设 %d 列' % len(col_wide))

    wb.save(path)
    wb.close()
    log('lang 归一化完成: %d 行, %d 语言 => %s' % (row_e - row_s + 1, len(langs), _rel_ui(path)))
    return langs, row_e - row_s + 1


# 统计文件
def encode_scui_lang_parser_txt(file, ws, langs, row_s, row_e, col_s, col_e):
    char_set = set()
    for idx_col in range(len(langs)):
        for idx_row in range(row_e - row_s + 1):
            text = _cell_text(ws.cell(row_s + idx_row, col_s + idx_col))
            char_set = char_set | set(text)
    char_set = sorted(list(char_set))
    print("char_set:\n", char_set)
    for item in char_set:
        file.write(item)


# 编写集成化源文件(bin: 列优先展开)
def encode_scui_lang_parser_bin(file, ws, langs, row_s, row_e, col_s, col_e):
    for idx_col in range(len(langs)):
        for idx_row in range(row_e - row_s + 1):
            c_bytes = _cell_text(ws.cell(row_s + idx_row, col_s + idx_col)).encode('utf-8')
            file.write(c_bytes)


# 编写集成化源文件
def encode_scui_lang_parser_c(file, ws, langs, row_s, row_e, col_s, col_e):
    n_row = row_e - row_s + 1
    n_col = len(langs)
    file.write('/* 本文件由 scui_pack_tools.exe 生成 */\n\n')
    file.write('#include "scui.h"\n\n')
    # 如果使用外源载入, 需要对应宏控制
    file.write('#if SCUI_LANG_PARSER_BIN_USE\n')
    # 提取所有外源依赖
    c_bytes_offset = 0
    c_bytes_length_max = 0
    for idx_col in range(n_col):
        for idx_row in range(n_row):
            c_bytes = _cell_text(ws.cell(row_s + idx_row, col_s + idx_col)).encode('utf-8')
            c_bytes_len = len(c_bytes)
            c_bytes_prefix = 'static const scui_lang_item_t scui_lang_parser_item_'
            c_bytes_format = '{:03d}_{:03d} = {{.offset = 0x{:08X}, .length = 0x{:08X},}};\n'
            file.write(c_bytes_prefix + c_bytes_format.format(idx_col, idx_row, c_bytes_offset, c_bytes_len))
            c_bytes_offset += c_bytes_len
            if c_bytes_length_max <= c_bytes_len:
                c_bytes_length_max = c_bytes_len
    file.write('\n')
    file.write('const void * const scui_lang_parser_table[%d * %d] = {\n' % (n_row, n_col))
    for idx_col in range(n_col):
        for idx_row in range(n_row):
            file.write('\t(void *)&scui_lang_parser_item_%03d_%03d,\n' % (idx_col, idx_row))
    file.write('};\n')
    file.write('#else\n')
    # 提取所有外源依赖
    # 字符串太特殊了, 直接用原生表形式处理
    file.write('const void * const scui_lang_parser_table[%d * %d] = {\n' % (n_row, n_col))
    for idx_col in range(n_col):
        for idx_row in range(n_row):
            # 多国语这里不能很好的对齐,因为不同语言的空格宽度不一致
            data = _cell_text(ws.cell(row_s + idx_row, col_s + idx_col))
            c_array = str([hex(c) for c in data.encode('utf-8')])
            c_array = c_array.replace('\'', '')
            c_array = c_array.replace('[', '{')
            c_array = c_array.replace(']', ', 0}')
            c_array = c_array.replace('{', '(const char []){')
            file.write('\t//--->>>%s\n\t%s,\n' % (data.replace('\n', '\\n'), c_array))
    file.write('};\n')
    file.write('#endif\n')
    # 填充函数定义或者声明
    file.write('\nstatic scui_lang_type_t scui_lang_type = 0;\n\n')
    file.write('/*@brief 获取多国语语言类型\n')
    file.write(' *@param type 语言类型编号\n */\n\n')
    file.write('void scui_lang_get(scui_lang_type_t *type)\n{\n')
    file.write('\tSCUI_ASSERT(type != NULL);\n\t*type = scui_lang_type;\n}\n\n')
    # 填充函数定义或者声明
    file.write('/*@brief 设置多国语语言类型\n')
    file.write(' *@param type 语言类型编号\n */\n\n')
    file.write('void scui_lang_set(scui_lang_type_t *type)\n{\n')
    file.write('\tSCUI_ASSERT(type != NULL);\n\tscui_lang_type = *type;\n\t\n')
    file.write('\tscui_event_define_absorb_none(event, SCUI_HANDLE_SYSTEM, false, ')
    file.write('scui_event_lang_change);\n')
    file.write('\tscui_event_notify(&event);\n}\n\n')
    # 填充函数定义或者声明
    file.write('/*@brief 多国语字符串转换\n')
    file.write(' *       需要同步拷贝使用\n')
    file.write(' *@param handle 字符串句柄\n')
    file.write(' *@param type   语言类型编号\n')
    file.write(' *@retval 字符串\n */\n')
    file.write('const char * scui_lang_str(scui_handle_t handle, scui_lang_type_t type)')
    file.write('\n{\n\tscui_handle_t string = handle;\n')
    file.write('\tswitch (type) {\n')
    file.write('\tdefault: string += type; break;\n')
    for item in SCUI_LANG_CUSTOM:
        file.write('\tcase scui_lang_type_%s: string += scui_lang_type; break;\n' % item)
    file.write('\t}\n\tstring -= scui_lang_ofs_num;\n\t\n')
    file.write('\t#if SCUI_LANG_PARSER_BIN_USE\n')
    file.write('\tstatic uint8_t scui_lang_str_buffer[0x%X + 1] = {0};\n' % c_bytes_length_max)
    file.write('\tconst scui_lang_item_t *item_utf8 = scui_handle_source(string);\n')
    file.write('\tSCUI_ASSERT(item_utf8->length < 0x%X + 1);\n' % c_bytes_length_max);
    file.write('\tSCUI_ASSERT(item_utf8->offset + item_utf8->length < 0x%X);\n' % c_bytes_offset);
    file.write('\tscui_lang_src_read(scui_lang_str_buffer, item_utf8->offset, item_utf8->length);\n')
    file.write('\tscui_lang_str_buffer[item_utf8->length] = \'\\0\';\n')
    file.write('\tconst char *str_utf8 = scui_lang_str_buffer;\n')
    file.write('\t#else\n')
    file.write('\tconst char *str_utf8 = scui_handle_source(string);\n')
    file.write('\t#endif\n')
    file.write('\treturn str_utf8;\n')
    file.write('}\n\n\n')


# 编写集成化头文件
def encode_scui_lang_parser_h(file, ws, langs, row_s, row_e, col_s, col_e, offset_name, offset_value):
    n_row = row_e - row_s + 1
    n_col = len(langs)
    file.write('#ifndef SCUI_LANG_PARSER_H\n')
    file.write('#define SCUI_LANG_PARSER_H\n\n')
    file.write('/* 本文件由 scui_pack_tools.exe 生成 */\n\n')
    # 编写头部索引(拼音区: 第2列; 编号区: 第1列)
    # 每区以 __OFFSET 给出该区起始偏移, 后续条目不再指定值(枚举自增)
    #   __OFFSET = SCUI_HANDLE_OFFSET_LANG, 则首条目 = OFFSET+1 = 0x5000
    #   (scui_ready.c 以 OFFSET+1 注册句柄表, 见 scui_lang_str 的偏移算式)
    # 列对齐: 名字列宽 = 两区最长名字(含 __OFFSET), 逗号后 \t + 注释
    abbr_name = []
    for idx, r in enumerate(range(row_s, row_e + 1)):
        name = _cell_text(ws.cell(r, SCUI_LANG_COL_ABBR + 1)).strip()
        if not name:
            name = 'SCUI_LANG_ABBR_0X%04X' % idx
        abbr_name.append(name)
    idx_name = ['SCUI_LANG_IDX_0X%04x' % i for i in range(row_e - row_s + 1)]
    abbr_ofs = 'SCUI_LANG_ABBR__OFFSET'
    idx_ofs  = 'SCUI_LANG_IDX__OFFSET'
    col_name = max([len(x) for x in abbr_name + idx_name + [abbr_ofs, idx_ofs]] + [0])
    file.write('typedef enum {\n')
    file.write('\t%s = %s,\n\t\n' % (offset_name, offset_value))
    file.write('\t/* 拼音区: */\n')
    file.write('\t%-*s = %s,\n' % (col_name, abbr_ofs, offset_name))
    for idx, r in enumerate(range(row_s, row_e + 1)):
        c_str = _cell_text(ws.cell(r, col_s)).replace('\n', '\\n')
        file.write('\t%-*s,\t/* %s */\n' % (col_name, abbr_name[idx], c_str))
    file.write('\t\n\t/* 编号区: */\n')
    file.write('\t%-*s = %s,\n' % (col_name, idx_ofs, offset_name))
    for idx, r in enumerate(range(row_s, row_e + 1)):
        c_str = _cell_text(ws.cell(r, col_s)).replace('\n', '\\n')
        file.write('\t%-*s,\t/* %s */\n' % (col_name, idx_name[idx], c_str))
    file.write('} scui_lang_parser_str_t;\n\n')
    # 如果使用外源载入, 需要对应宏控制
    file.write('#if SCUI_LANG_PARSER_BIN_USE\n')
    file.write('typedef struct {\n')
    file.write('\tuintptr_t offset;\n')
    file.write('\tuintptr_t length;\n')
    file.write('} scui_lang_item_t;\n')
    file.write('#endif\n\n')
    file.write('extern const void * const scui_lang_parser_table[%d * %d];\n\n' % (n_row, n_col))
    # 编写头部索引
    file.write('typedef enum {\n')
    file.write('\tscui_lang_str_num = %s,\n' % str(n_row))
    file.write('\tscui_lang_ofs_num = 1,\n\t\n')
    file.write('\tscui_lang_type_num\t\t= %s,\n' % str(n_col))
    for idx, item in enumerate(SCUI_LANG_CUSTOM):
        file.write('\tscui_lang_type_%s\t= SCUI_HANDLE_SYSTEM - %d,\n' % (item, idx + 1))
    for idx, item in enumerate(langs):
        file.write('\tscui_lang_type_%s\t\t= %s * %s + scui_lang_ofs_num,\n' % (item, 'scui_lang_str_num', str(idx)))
    file.write('} scui_lang_type_t;\n\n')
    # 填充函数定义或者声明
    file.write('/*@brief 获取多国语语言类型\n')
    file.write(' *@param type 语言类型编号\n */\n')
    file.write('void scui_lang_get(scui_lang_type_t *type);\n\n')
    # 填充函数定义或者声明
    file.write('/*@brief 设置多国语语言类型\n')
    file.write(' *@param type 语言类型编号\n */\n')
    file.write('void scui_lang_set(scui_lang_type_t *type);\n\n')
    # 填充函数定义或者声明
    file.write('/*@brief 多国语字符串转换\n')
    file.write(' *       需要同步拷贝使用\n')
    file.write(' *@param handle 字符串句柄\n')
    file.write(' *@param type   语言类型编号\n')
    file.write(' *@retval 字符串\n */\n')
    file.write('const char * scui_lang_str(scui_handle_t handle, scui_lang_type_t type);\n\n')
    file.write('#endif\n')


# 打印重定向到文件
class ScuiRedirectPrint(object):
    def __init__(self, stream=sys.stdout, path='.', file='log.txt'):
        if not os.path.exists(path):
            os.makedirs(path)
        self.log = open(os.path.join(path, file), mode='w+', encoding='utf-8')
        self.terminal = stream

    def write(self, message):
        self.log.write(message)
        self.terminal.write(message)

    def flush(self):
        pass


# 启用集成化资源总表生成
def encode_scui_lang_parser():
    # 参数列表:文件根目录,
    if len(sys.argv) != 3:
        print('argv list not match')
        return
    src_path = sys.argv[1]
    dst_path = sys.argv[2]
    # 获得文件处理根路径
    if not os.path.exists(src_path):
        print('src path is not exist')
        return
    if not os.path.exists(dst_path):
        print('dst path is not exist')
        return
    print('src path:', _rel_ui(src_path))
    print('dst path:', _rel_ui(dst_path))
    # 句柄偏移: 前端"设置句柄"json 覆盖配置默认
    offset_name, offset_value = _handle_offset('lang', SCUI_LANG_OFFSET_NAME, SCUI_LANG_OFFSET_VALUE)
    # 打包默认执行一次归一化(直接复写 xlsx)
    normalize_scui_lang_sheet(src_path, log=print)
    # 以落盘结果为准重新打开工作簿
    xlsx_file = openpyxl.load_workbook(_sheet_path(src_path), read_only=False)
    xlsx_sheet = _sheet_pick(xlsx_file)
    langs, row_s, row_e, col_s, col_e = _sheet_layout(xlsx_sheet)
    if not langs:
        print('sheet language list is empty(head row)')
        xlsx_file.close()
        return
    # 开启四个文件
    scui_lang_parser_h = open(os.path.join(dst_path, 'scui_res_lang.h'), mode='w', encoding='utf-8')
    scui_lang_parser_c = open(os.path.join(dst_path, 'scui_res_lang.c'), mode='w', encoding='utf-8')
    scui_lang_parser_bin = open(os.path.join(dst_path, 'scui_res_lang.bin'), mode='wb')
    scui_lang_parser_txt = open(os.path.join(dst_path, 'scui_res_lang.txt'), mode='w', encoding='utf-8')
    # 解析
    encode_scui_lang_parser_h(scui_lang_parser_h, xlsx_sheet, langs, row_s, row_e, col_s, col_e,
                              offset_name, offset_value)
    encode_scui_lang_parser_c(scui_lang_parser_c, xlsx_sheet, langs, row_s, row_e, col_s, col_e)
    encode_scui_lang_parser_bin(scui_lang_parser_bin, xlsx_sheet, langs, row_s, row_e, col_s, col_e)
    encode_scui_lang_parser_txt(scui_lang_parser_txt, xlsx_sheet, langs, row_s, row_e, col_s, col_e)
    # 关闭四个文件
    scui_lang_parser_h.close()
    scui_lang_parser_c.close()
    scui_lang_parser_bin.close()
    scui_lang_parser_txt.close()
    # 关闭工作簿
    xlsx_file.close()
    print('lang rows:%d cols:%d' % (row_e - row_s + 1, len(langs)))


if __name__ == '__main__':
    try:
        encode_scui_lang_parser()
    except Exception as e:
        print(e)
    print('scui lang parser finish')
    input('请按任意键退出...')
